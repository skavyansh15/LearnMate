"""Builds the DEMO course in demo_course/: subtitles, PDF, PPTX, silent video, interview questions and course.json.
Needs: pip install reportlab pillow python-pptx  and ffmpeg.  Run: python tools/make_demo.py   (not needed for your own courses)"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import ingest
import store

DATA_DIR = VIDEO_DIR = Path(__file__).resolve().parent.parent / "demo_course"
SRT_FILE, PDF_FILE, PPTX_FILE = DATA_DIR / "lecture.srt", DATA_DIR / "lecture_notes.pdf", DATA_DIR / "lecture_slides.pptx"
MAPPING_FILE, VIDEO_FILE = DATA_DIR / "mapping.json", "lecture.mp4"

SEGMENTS = [  # (title, code, sentence 1, sentence 2)
    ("What is a list?", 'fruits = ["apple", "banana", "cherry"]',
     "Welcome to this short lesson on Python lists.",
     "A list stores many values in one variable, written inside square brackets and separated by commas."),
    ("Indexing", 'fruits[0]  # "apple"      fruits[1]  # "banana"',
     "Every item in a list has a position called an index.",
     "Python starts counting at zero, so fruits of zero is apple, and fruits of one is banana."),
    ("Negative indexing", 'fruits[-1]  # "cherry"',
     "You can also count from the end using negative indexes.",
     "Fruits of minus one gives the last item, which is cherry, without needing to know the length of the list."),
    ("Adding items", 'fruits.append("mango")   fruits.insert(0, "kiwi")',
     "To add an item at the end, use the append method, for example fruits dot append mango.",
     "If you need a specific position, insert takes an index and a value, so fruits dot insert zero kiwi puts kiwi at the front."),
    ("Removing items", 'fruits.remove("banana")   fruits.pop()',
     "To delete a value by name use remove, for example fruits dot remove banana.",
     "The pop method removes an item by position and also returns it, and with no argument it takes the last item."),
    ("Slicing", "fruits[1:3]  # items at index 1 and 2",
     "Slicing gives you part of a list.",
     "Writing fruits one colon three starts at index one and stops before index three, so the end index is not included."),
    ("Looping", "for fruit in fruits:\n    print(fruit)",
     "To visit every item, use a for loop: for fruit in fruits, then print the fruit.",
     "The loop variable takes each value in the list, one at a time, from first to last."),
    ("List comprehension", "squares = [x * x for x in range(5)]",
     "A list comprehension builds a new list in one line.",
     "For example, squares equals x times x for x in range five, which gives zero, one, four, nine, sixteen."),
]
PDF_PAGES = [
    ("Lists and indexing", 'A list stores several values in one variable. Write it in square brackets: fruits = ["apple", "banana", "cherry"]. Each item has an index. Indexes start at 0, so fruits[0] is "apple" and fruits[1] is "banana". Negative indexes count from the end: fruits[-1] is "cherry".'),
    ("Changing a list", 'append adds an item at the end: fruits.append("mango"). insert(index, value) adds at a position: fruits.insert(0, "kiwi"). remove(value) deletes the first matching value. pop(index) removes by position and returns the item; pop() with no argument removes the last item.'),
    ("Slicing", "fruits[1:3] returns the items at index 1 and 2. The start index is included and the end index is not. Leaving out a number means from the start or to the end, for example fruits[:2] and fruits[1:]."),
    ("Loops and comprehensions", "A for loop visits every item: for fruit in fruits: print(fruit). A list comprehension builds a new list in one line: squares = [x * x for x in range(5)] gives [0, 1, 4, 9, 16]."),
]
SLIDES = [
    ("What is a list?", ["Many values in one variable", "Square brackets, comma separated"]),
    ("Indexing", ["Each item has an index", "Counting starts at 0", "fruits[0] is the first item"]),
    ("Negative indexing", ["Count from the end", "fruits[-1] is the last item"]),
    ("Adding and removing", ["append(x): add at the end", "insert(i, x): add at position i", "remove(x): delete by value", "pop(): remove and return the last item"]),
    ("Slicing", ["fruits[1:3]", "Start included, end not included"]),
    ("Loops and comprehensions", ["for item in list:", "squares = [x * x for x in range(5)]"]),
]
# page/slide -> video second when it unlocks (rule: when the LAST topic on it has started)
MAPPING = {"pdf": {"1": 60, "2": 120, "3": 150, "4": 210},
           "slide": {"1": 0, "2": 30, "3": 60, "4": 120, "5": 150, "6": 210}}


def srt_time(s):
    return f"00:{int(s) // 60:02d}:{int(s) % 60:02d},000"


def make_srt():
    out, n = [], 1
    for i, seg in enumerate(SEGMENTS):
        for j, sentence in enumerate(seg[2:]):
            a = i * 30 + j * 15
            out.append(f"{n}\n{srt_time(a)} --> {srt_time(a + 15)}\n{sentence}\n")
            n += 1
    SRT_FILE.write_text("\n".join(out), encoding="utf-8")


def make_pdf():
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    st = getSampleStyleSheet()
    flow = []
    for i, (title, body) in enumerate(PDF_PAGES):
        flow += [Paragraph(title, st["Heading1"]), Spacer(1, 12), Paragraph(body.replace("&", "&amp;"), st["BodyText"])]
        if i < len(PDF_PAGES) - 1:
            flow.append(PageBreak())
    SimpleDocTemplate(str(PDF_FILE)).build(flow)


def make_pptx():
    from pptx import Presentation
    prs = Presentation()
    for title, bullets in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
    prs.save(PPTX_FILE)


def make_video():
    from PIL import Image, ImageDraw, ImageFont
    f = "/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf"
    try:
        big, mono, small = (ImageFont.truetype(f % "-Bold", 44), ImageFont.truetype(f % "", 28), ImageFont.truetype(f % "", 22))
    except OSError:
        big = mono = small = ImageFont.load_default()
    tmp = Path(tempfile.mkdtemp())
    lines = []
    for i, (title, code, s1, s2) in enumerate(SEGMENTS):
        for j, caption in enumerate((s1, s2)):
            img = Image.new("RGB", (960, 540), "#14172a")
            d = ImageDraw.Draw(img)
            d.text((50, 50), f"Python lists - {i + 1}/8", font=small, fill="#9a9db2")
            d.text((50, 110), title, font=big, fill="#ffb020")
            d.multiline_text((50, 220), code, font=mono, fill="#e8e9f0", spacing=8)
            words, row, rows = caption.split(), "", []
            for w in words:
                if len(row + w) > 62:
                    rows.append(row)
                    row = ""
                row += w + " "
            rows.append(row)
            d.multiline_text((50, 420), "\n".join(rows), font=small, fill="#2dd4bf", spacing=6)
            p = tmp / f"f{i}_{j}.png"
            img.save(p)
            lines += [f"file '{p}'", "duration 15"]
    lines.append(lines[-2])  # concat demuxer needs the last file repeated
    (tmp / "list.txt").write_text("\n".join(lines))
    VIDEO_DIR.mkdir(exist_ok=True)
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(tmp / "list.txt"),
                    "-vf", "fps=10,format=yuv420p", "-c:v", "libx264", "-g", "10", "-movflags", "+faststart",
                    str(VIDEO_DIR / VIDEO_FILE)], check=True)



BANK = [
    {"question": "What is a Python list and how do you create one?", "level": "easy", "key_points": [
        {"name": "stores many values in one variable", "terms": ["many", "multiple", "several", "collection", "values", "items"]},
        {"name": "square brackets", "terms": ["bracket", "[]", "["]},
        {"name": "items separated by commas", "terms": ["comma"]}]},
    {"question": "How does indexing work in a list?", "level": "easy", "key_points": [
        {"name": "counting starts at zero", "terms": ["zero", "0"]},
        {"name": "index is the position of an item", "terms": ["position", "index"]},
        {"name": "negative index counts from the end", "terms": ["negative", "-1", "end", "last"]}]},
    {"question": "How do you add items to a list?", "level": "easy", "key_points": [
        {"name": "append adds at the end", "terms": ["append"]},
        {"name": "insert adds at a chosen position", "terms": ["insert"]},
        {"name": "position or index is given to insert", "terms": ["position", "index"]}]},
    {"question": "Explain the difference between remove and pop.", "level": "hard", "key_points": [
        {"name": "remove deletes by value", "terms": ["value", "name"]},
        {"name": "pop removes by position", "terms": ["position", "index"]},
        {"name": "pop returns the removed item", "terms": ["return", "returns", "gives back"]}]},
    {"question": "What does fruits[1:3] return, and why?", "level": "hard", "key_points": [
        {"name": "items at index 1 and 2", "terms": ["index 1", "1 and 2", "second and third"]},
        {"name": "start index is included", "terms": ["included", "inclusive", "includes"]},
        {"name": "end index is not included", "terms": ["not included", "excluded", "exclusive", "stops before", "up to but"]}]},
    {"question": "What is a list comprehension and when would you use it?", "level": "hard", "key_points": [
        {"name": "builds a new list", "terms": ["new list", "build", "create"]},
        {"name": "in one line", "terms": ["one line", "single line", "short", "concise"]},
        {"name": "uses an expression with a for loop inside brackets", "terms": ["for", "loop", "expression"]}]},
]


def build_course():
    mapping = json.loads(MAPPING_FILE.read_text())
    cues = ingest.parse_subtitles(SRT_FILE.read_text())
    chunks = ingest.chunk_cues(cues)
    pdf_pages = [(i + 1, f"{t}\n{b}") for i, (t, b) in enumerate(PDF_PAGES)]  # same text that was written into the PDF
    slides = ingest.read_pptx(PPTX_FILE.read_bytes())                          # parsed from the real PPTX
    chunks += ingest.chunks_from_pages("pdf", pdf_pages, "lecture_notes.pdf", mapping["pdf"])
    chunks += ingest.chunks_from_pages("slide", slides, "lecture_slides.pptx", mapping["slide"])
    for i, c in enumerate(chunks):
        c["id"] = f"c{i}"
    course = {"id": store.DEMO_ID, "title": "Python Lists (demo course)", "created": 0, "has_video": True,
              "video_url": "/media/demo.mp4", "cues": cues, "chunks": chunks,
              "documents": [{"id": "d1", "name": "lecture_notes.pdf", "kind": "pdf", "pages": len(pdf_pages)},
                            {"id": "d2", "name": "lecture_slides.pptx", "kind": "slide", "pages": len(slides)}]}
    (DATA_DIR / "course.json").write_text(json.dumps(course, ensure_ascii=False))
    (DATA_DIR / "interview_bank.json").write_text(json.dumps(BANK, indent=1))
    print(f"course.json: {len(chunks)} chunks")


if __name__ == "__main__":
    make_srt(); make_pdf(); make_pptx()
    MAPPING_FILE.write_text(json.dumps(MAPPING, indent=2))
    make_video()
    build_course()
