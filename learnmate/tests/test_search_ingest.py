import io
import sys
import types
import unittest
from unittest import mock

from helpers import Base, config
import ingest
import search
from errors import UserError

SRT = "1\n00:00:00,000 --> 00:00:15,000\nHello there friend.\n\n2\n00:00:15,000 --> 00:00:30,000\nSecond line here.\n"
VTT = "WEBVTT\n\n00:00.000 --> 00:15.000\nFirst <b>cue</b>\n\n00:15.000 --> 00:31.500\nSecond cue\n"


class SearchTests(unittest.TestCase):
    def test_stemming_matches_word_forms(self):
        self.assertEqual(search.tokens("slicing"), search.tokens("slice"))
        self.assertEqual(search.tokens("indexes"), search.tokens("index"))

    def test_no_evidence_returns_nothing(self):
        chunks = [{"text": "A list stores many values."}]
        self.assertEqual(search.search(chunks, "what is a dictionary"), [])

    def test_rare_missing_word_blocks_weak_match(self):
        chunks = [{"text": "Every item in a list has an index."}, {"text": "A list stores values."}]
        self.assertEqual(search.search(chunks, "how do I append an item to a list"), [])

    def test_finds_best_chunk(self):
        chunks = [{"text": "Slicing takes part of a list."}, {"text": "Loops visit every item."}]
        self.assertIs(search.search(chunks, "how does slicing work")[0], chunks[0])


class SubtitleTests(unittest.TestCase):
    def test_srt(self):
        cues = ingest.parse_subtitles(SRT)
        self.assertEqual([(c["start"], c["end"]) for c in cues], [(0, 15), (15, 30)])

    def test_vtt_without_hours_and_with_tags(self):
        cues = ingest.parse_subtitles(VTT)
        self.assertEqual(cues[0]["text"], "First cue")
        self.assertEqual(cues[1]["end"], 31.5)

    def test_bad_transcript_gives_helpful_error(self):
        with self.assertRaises(UserError) as cm:
            ingest.parse_subtitles("just some text")
        self.assertIn(".srt", cm.exception.hint)

    def test_chunks_are_30_to_60_seconds_and_unlock_at_start(self):
        cues = [{"start": i * 15, "end": i * 15 + 15, "text": f"line {i}"} for i in range(16)]
        chunks = ingest.chunk_cues(cues)
        self.assertEqual(len(chunks), 8)
        for c in chunks:
            self.assertTrue(30 <= c["end"] - c["start"] <= 60)
            self.assertEqual(c["unlock"], c["start"])


class DocumentTests(Base):
    def test_long_page_split_keeps_page_number(self):
        text = " ".join(f"Sentence number {i} is here." for i in range(120))
        chunks = ingest.chunks_from_pages("pdf", [(7, text)], "book.pdf")
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(c["ref"] == 7 and c["label"] == "book.pdf p.7" for c in chunks))

    def test_unmapped_pages_are_skipped_not_guessed(self):
        chunks = ingest.chunks_from_pages("slide", [(1, "a b c"), (2, "d e f")], "d.pptx", {"1": 0})
        self.assertEqual([c["ref"] for c in chunks], [1])

    def test_txt_line_ranges_are_exact(self):
        text = "\n".join(f"line {i} " + "x" * 100 for i in range(1, 31))
        chunks = ingest.chunks_from_txt(text, "n.txt", max_chars=500)
        self.assertEqual(chunks[0]["label"], "n.txt lines 1-5")
        self.assertEqual(chunks[-1]["label"].split("-")[-1], "30")

    def test_validation_errors(self):
        cases = [("a.exe", b"MZ", "bad_type"), ("a.txt", b"", "empty_file"), ("a.pdf", b"not a pdf", "bad_pdf"),
                 ("a.pptx", b"not a zip", "bad_pptx"), ("a.txt", b"abc\x00def", "bad_txt"), ("a.txt", b"\xff\xfe\x00bad", "bad_txt"),
                 ("a.txt", b"\xc3\x28 broken", "bad_txt")]
        for name, data, code in cases:
            with self.subTest(name=name, code=code):
                with self.assertRaises(UserError) as cm:
                    ingest.extract_document(name, data)
                self.assertEqual(cm.exception.code, code)
                self.assertTrue(cm.exception.hint or cm.exception.message)

    def test_too_big(self):
        with mock.patch.object(config, "MAX_UPLOAD_MB", 0):
            with self.assertRaises(UserError) as cm:
                ingest.extract_document("a.txt", b"hello world")
        self.assertEqual(cm.exception.code, "too_big")

    def test_real_pptx_has_exact_slide_numbers(self):
        from pptx import Presentation
        prs = Presentation()
        for title in ("Intro slide", "Second slide"):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = title
            s.placeholders[1].text_frame.text = f"Body of {title}"
        buf = io.BytesIO()
        prs.save(buf)
        kind, chunks, count = ingest.extract_document("deck.pptx", buf.getvalue())
        self.assertEqual((kind, count), ("slide", 2))
        self.assertEqual([c["label"] for c in chunks], ["deck.pptx slide 1", "deck.pptx slide 2"])
        self.assertIn("Body of Second slide", chunks[1]["text"])

    def test_pdf_path_with_fake_pymupdf(self):
        """Checks OUR code around PyMuPDF (page numbers, labels). The real library is tested in test_real_pymupdf."""
        class Page:
            def __init__(self, t): self.t = t
            def get_text(self): return self.t

        class Doc:
            page_count = 2
            def __iter__(self): return iter([Page("First page text here."), Page("Second page text here.")])

        fake = types.SimpleNamespace(open=lambda **kw: Doc())
        with mock.patch.dict(sys.modules, {"fitz": fake}):
            kind, chunks, count = ingest.extract_document("n.pdf", b"%PDF-1.4 fake")
        self.assertEqual([c["label"] for c in chunks], ["n.pdf p.1", "n.pdf p.2"])

    @unittest.skipUnless(__import__("importlib").util.find_spec("fitz"), "PyMuPDF not installed here")
    def test_real_pymupdf(self):
        data = (config.DEMO_DIR / "lecture_notes.pdf").read_bytes()
        kind, chunks, count = ingest.extract_document("lecture_notes.pdf", data)
        self.assertEqual(count, 4)
        self.assertIn("append", next(c for c in chunks if c["ref"] == 2)["text"])


if __name__ == "__main__":
    unittest.main()
